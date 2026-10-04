"""Word / PPT CRUD smoke tests (disk)."""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from pptx import Presentation

from adapters.word.word_adapter import WordAdapter
from adapters.ppt.ppt_adapter import PptAdapter
from session.edit_engine import _op_to_preview
from session.session_manager import Session


def test_word_table_and_insert():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "w.docx")
        doc = Document()
        doc.add_paragraph("Intro")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Q"
        table.cell(0, 1).text = "A"
        doc.save(path)

        adapter = WordAdapter(path)
        p = adapter.write_table_cell(0, 0, 1, "Answer")
        p.apply()
        adapter.reload()
        assert "Answer" in adapter._doc.tables[0].cell(0, 1).text

        p2 = adapter.insert_paragraph(0, "Title")
        p2.apply()
        adapter.reload()
        assert adapter.read_paragraph(0) == "Title"


def test_ppt_slide_ops():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "p.pptx")
        prs = Presentation()
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "Hello"
        prs.save(path)

        adapter = PptAdapter(path)
        preview = adapter.set_slide_text(1, "Updated")
        preview.apply()
        adapter.reload()
        assert "Updated" in adapter.read_all()

        adapter.add_slide("New").apply()
        adapter.reload()
        assert len(list(adapter._prs.slides)) >= 2


def test_edit_engine_word_op():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "w.docx")
        doc = Document()
        doc.add_paragraph("One")
        doc.save(path)
        session = Session("t")
        session.merge_file(path)
        preview = _op_to_preview(
            session,
            {"op": "append_paragraph", "target": "w.docx", "value": "Two"},
        )
        preview.apply()
        adapter = WordAdapter(path)
        assert "Two" in adapter.read_all()


if __name__ == "__main__":
    test_word_table_and_insert()
    test_ppt_slide_ops()
    test_edit_engine_word_op()
    print("Word/PPT tests passed.")
