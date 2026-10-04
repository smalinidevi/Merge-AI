import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from adapters.excel.excel_adapter import ExcelAdapter
from session.session_manager import SessionManager


def make_sample_xlsx(path, sheet_data):
    """sheet_data: {sheet_name: [[row1...], [row2...]]}"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheet_data.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)


def test_read_cell():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "a.xlsx")
        make_sample_xlsx(path, {"Sheet1": [["Revenue", 100], ["Cost", 40]]})
        adapter = ExcelAdapter(path)
        assert adapter.read_cell("A1", "Sheet1") == "Revenue"
        assert adapter.read_cell("B1", "Sheet1") == 100


def test_write_cell_confirm_first():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "a.xlsx")
        make_sample_xlsx(path, {"Sheet1": [["Revenue", 100]]})
        adapter = ExcelAdapter(path)
        preview = adapter.write_cell("B1", 200, "Sheet1")
        # not applied yet
        assert adapter.read_cell("B1", "Sheet1") == 100
        preview.apply()
        # reload from disk to confirm it actually persisted
        adapter2 = ExcelAdapter(path)
        assert adapter2.read_cell("B1", "Sheet1") == 200


def test_append_and_delete_row():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "a.xlsx")
        make_sample_xlsx(path, {"Sheet1": [["Revenue", 100]]})
        adapter = ExcelAdapter(path)
        adapter.append_row(["Cost", 40], "Sheet1").apply()

        adapter2 = ExcelAdapter(path)
        assert adapter2.read_cell("A2", "Sheet1") == "Cost"

        adapter2.delete_row(2, "Sheet1").apply()
        adapter3 = ExcelAdapter(path)
        assert adapter3.read_cell("A2", "Sheet1") is None


def test_session_multi_file_search_and_banner():
    with tempfile.TemporaryDirectory() as tmp:
        path_a = os.path.join(tmp, "a.xlsx")
        path_b = os.path.join(tmp, "b.xlsx")
        make_sample_xlsx(path_a, {"Sheet1": [["Revenue", 100]]})
        make_sample_xlsx(path_b, {"Sheet1": [["Total", 100]]})

        manager = SessionManager()
        session = manager.new_session()
        session.merge_file(path_a)
        session.merge_file(path_b)

        banner = session.banner_text()
        assert "a.xlsx" in banner and "b.xlsx" in banner

        results = session.search_all(100)
        assert "a.xlsx" in results and "b.xlsx" in results


def test_session_detects_stale_source():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "a.xlsx")
        make_sample_xlsx(path, {"Sheet1": [["x", 1]]})
        manager = SessionManager()
        session = manager.new_session()
        session.merge_file(path)
        os.remove(path)
        banner = session.banner_text()
        assert "No longer available" in banner


def test_create_rename_delete_sheet():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "a.xlsx")
        make_sample_xlsx(path, {"Sheet1": [["x", 1]]})
        adapter = ExcelAdapter(path)
        adapter.create_sheet("NewData").apply()
        adapter2 = ExcelAdapter(path)
        assert "NewData" in adapter2.list_sheets()
        adapter2.rename_sheet("NewData", "Report").apply()
        adapter3 = ExcelAdapter(path)
        assert "Report" in adapter3.list_sheets()
        assert "NewData" not in adapter3.list_sheets()
        adapter3.copy_sheet("Sheet1", "Sheet1Copy").apply()
        adapter4 = ExcelAdapter(path)
        assert "Sheet1Copy" in adapter4.list_sheets()
        adapter4.delete_sheet("Report").apply()
        adapter5 = ExcelAdapter(path)
        assert "Report" not in adapter5.list_sheets()


def test_insert_column_and_write_range():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "a.xlsx")
        make_sample_xlsx(path, {"Sheet1": [["A", "B"], [1, 2]]})
        adapter = ExcelAdapter(path)
        adapter.insert_column("B", "Sheet1").apply()
        adapter2 = ExcelAdapter(path)
        # After insert at B, old B shifts to C
        assert adapter2.read_cell("C1", "Sheet1") == "B"
        adapter2.write_range("A3", [["x", "y", "z"]], "Sheet1").apply()
        adapter3 = ExcelAdapter(path)
        assert adapter3.read_cell("A3", "Sheet1") == "x"
        assert adapter3.read_cell("C3", "Sheet1") == "z"


if __name__ == "__main__":
    test_read_cell()
    test_write_cell_confirm_first()
    test_append_and_delete_row()
    test_session_multi_file_search_and_banner()
    test_session_detects_stale_source()
    test_create_rename_delete_sheet()
    test_insert_column_and_write_range()
    print("All tests passed.")
