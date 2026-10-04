"""Edit-engine mapping tests (no Azure / no overlay)."""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook
from docx import Document

from session.session_manager import Session
from session.edit_engine import (
    _op_to_preview,
    apply_previews,
    looks_like_edit,
    reload_writable_adapters,
    writable_sources,
)


def test_looks_like_edit():
    assert looks_like_edit("Set cell A1 to 42")
    assert looks_like_edit("delete row 3")
    assert not looks_like_edit("What is the total?")


def test_excel_write_via_op():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "ledger.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Old"
        wb.save(path)

        session = Session("t")
        session.merge_file(path)
        assert len(writable_sources(session)) == 1

        preview = _op_to_preview(
            session,
            {
                "op": "write_cell",
                "target": "ledger.xlsx",
                "sheet": "Sheet1",
                "cell": "A1",
                "value": "New",
            },
        )
        apply_previews([preview])
        reload_writable_adapters(session)
        adapter = session.sources[os.path.abspath(path)].adapter
        assert adapter.read_cell("A1", "Sheet1") == "New"


def test_word_append_via_op():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "notes.docx")
        Document().save(path)
        # python-docx empty doc still has one empty paragraph
        session = Session("t")
        session.merge_file(path)
        preview = _op_to_preview(
            session,
            {"op": "append_paragraph", "target": "notes.docx", "value": "Hello"},
        )
        apply_previews([preview])
        reload_writable_adapters(session)
        adapter = next(iter(session.sources.values())).adapter
        assert "Hello" in adapter.read_all()


def test_text_replace_via_op():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "note.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write("alpha beta gamma")
        session = Session("t")
        session.merge_file(path)
        preview = _op_to_preview(
            session,
            {
                "op": "replace_text",
                "target": "note.txt",
                "old": "beta",
                "new": "BETA",
            },
        )
        apply_previews([preview])
        with open(path, encoding="utf-8") as f:
            assert f.read() == "alpha BETA gamma"


if __name__ == "__main__":
    test_looks_like_edit()
    test_excel_write_via_op()
    test_word_append_via_op()
    test_text_replace_via_op()
    print("All edit_engine tests passed.")
