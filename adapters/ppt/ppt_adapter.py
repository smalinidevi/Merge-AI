"""
PowerPoint adapter (.pptx): read slide text + confirm-first replace/append.
"""

from __future__ import annotations

import os
from typing import Any

from pptx import Presentation

from adapters.base_adapter import BaseAdapter
from adapters.common.change_preview import ChangePreview


class PptAdapter(BaseAdapter):
    kind = "ppt"
    can_write = True

    def __init__(self, file_path: str):
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"PowerPoint file not found: {file_path}")
        super().__init__(os.path.abspath(file_path))
        self.file_path = self.source
        self._prs = Presentation(self.file_path)

    def reload(self) -> None:
        self._prs = Presentation(self.file_path)

    def read_all(self, max_chars: int = 20000) -> str:
        chunks = []
        for i, slide in enumerate(self._prs.slides, start=1):
            chunks.append(f"--- Slide {i} ---")
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text_frame.text.strip()
                    if text:
                        chunks.append(text)
        full = "\n".join(chunks)
        if len(full) > max_chars:
            return full[:max_chars] + f"\n... (truncated, {len(full) - max_chars} more chars)"
        return full

    def search(self, value: Any) -> list[str]:
        needle = str(value).lower()
        hits = []
        for i, slide in enumerate(self._prs.slides, start=1):
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    if needle in shape.text_frame.text.lower():
                        hits.append(f"slide {i}")
                        break
        return hits

    def replace_text(self, old: str, new: str) -> ChangePreview:
        matches = 0
        for slide in self._prs.slides:
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if old in run.text:
                            matches += run.text.count(old)

        if matches == 0:
            raise ValueError(f"Text not found in presentation: {old!r}")

        def _apply():
            for slide in self._prs.slides:
                for shape in slide.shapes:
                    if not getattr(shape, "has_text_frame", False):
                        continue
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new)
            self._prs.save(self.file_path)

        return ChangePreview(
            description="replace_text",
            source=self.file_path,
            location=f"{matches} run(s)",
            old_value=old,
            new_value=new,
            _apply_fn=_apply,
        )

    def _slide(self, slide_index: int):
        """1-based slide index."""
        slides = list(self._prs.slides)
        if slide_index < 1 or slide_index > len(slides):
            raise IndexError(f"slide {slide_index} out of range (1-{len(slides)})")
        return slides[slide_index - 1]

    def set_slide_text(self, slide: int, text: str) -> ChangePreview:
        """Replace text in the first text-bearing shape on the slide (1-based)."""
        s = self._slide(int(slide))
        target = None
        old = ""
        for shape in s.shapes:
            if getattr(shape, "has_text_frame", False):
                target = shape
                old = shape.text_frame.text
                break
        if target is None:
            raise ValueError(f"No text shape on slide {slide}")

        def _apply():
            tf = target.text_frame
            if tf.paragraphs:
                tf.paragraphs[0].text = str(text)
                for para in tf.paragraphs[1:]:
                    for run in para.runs:
                        run.text = ""
            else:
                tf.text = str(text)
            self._prs.save(self.file_path)

        return ChangePreview(
            description="set_slide_text",
            source=self.file_path,
            location=f"slide {slide}",
            old_value=old,
            new_value=text,
            _apply_fn=_apply,
        )

    def append_slide_text(self, slide: int, text: str) -> ChangePreview:
        s = self._slide(int(slide))
        target = None
        old = ""
        for shape in s.shapes:
            if getattr(shape, "has_text_frame", False):
                target = shape
                old = shape.text_frame.text
                break
        if target is None:
            raise ValueError(f"No text shape on slide {slide}")
        new_val = (old + ("\n" if old.strip() else "") + str(text)).strip()

        def _apply():
            p = target.text_frame.add_paragraph()
            p.text = str(text)
            self._prs.save(self.file_path)

        return ChangePreview(
            description="append_slide_text",
            source=self.file_path,
            location=f"slide {slide}",
            old_value=old,
            new_value=new_val,
            _apply_fn=_apply,
        )

    def add_slide(self, title: str | None = None) -> ChangePreview:
        layout = self._prs.slide_layouts[0]
        new_index = len(self._prs.slides) + 1

        def _apply():
            slide = self._prs.slides.add_slide(layout)
            if title:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        shape.text_frame.paragraphs[0].text = str(title)
                        break
            self._prs.save(self.file_path)

        return ChangePreview(
            description="add_slide",
            source=self.file_path,
            location=f"slide {new_index}",
            old_value=None,
            new_value=title or "(blank)",
            _apply_fn=_apply,
        )

    def delete_slide(self, slide: int) -> ChangePreview:
        idx = int(slide) - 1
        slides = list(self._prs.slides)
        if idx < 0 or idx >= len(slides):
            raise IndexError(f"slide {slide} out of range")
        # capture short preview of text
        old_bits = []
        for shape in slides[idx].shapes:
            if getattr(shape, "has_text_frame", False):
                t = shape.text_frame.text.strip()
                if t:
                    old_bits.append(t[:80])

        def _apply():
            rId = self._prs.slides._sldIdLst[idx].rId
            self._prs.part.drop_rel(rId)
            del self._prs.slides._sldIdLst[idx]
            self._prs.save(self.file_path)
            self.reload()

        return ChangePreview(
            description="delete_slide",
            source=self.file_path,
            location=f"slide {slide}",
            old_value="; ".join(old_bits) or "(empty)",
            new_value=None,
            _apply_fn=_apply,
        )
