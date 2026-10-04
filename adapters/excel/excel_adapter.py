"""
Excel adapter: full CRUD on a single .xlsx workbook, on top of BaseAdapter.

Supports cells, rows, columns, ranges, and sheet create/rename/delete/copy.
All writes are confirm-first via ChangePreview.
"""

from __future__ import annotations

import os
from typing import Any, Optional

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from adapters.base_adapter import BaseAdapter
from adapters.common.change_preview import ChangePreview


def _col_to_index(col: Any) -> int:
    """1-based column index from int or letter (A, B, AA…)."""
    if isinstance(col, int):
        if col < 1:
            raise ValueError("column index must be >= 1")
        return col
    text = str(col).strip().upper()
    if text.isdigit():
        idx = int(text)
        if idx < 1:
            raise ValueError("column index must be >= 1")
        return idx
    return column_index_from_string(text)


class ExcelAdapter(BaseAdapter):
    kind = "excel"
    can_write = True

    def __init__(self, file_path: str):
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"Excel file not found: {file_path}")
        super().__init__(os.path.abspath(file_path))
        self.file_path = self.source
        self._wb = openpyxl.load_workbook(self.file_path, data_only=False)

    # ---------- introspection ----------

    def list_sheets(self) -> list[str]:
        return list(self._wb.sheetnames)

    def _sheet(self, sheet_name: Optional[str] = None) -> Worksheet:
        if sheet_name is None:
            return self._wb.active
        if sheet_name not in self._wb.sheetnames:
            raise KeyError(f"Sheet '{sheet_name}' not found in {self.file_path}")
        return self._wb[sheet_name]

    def _save(self) -> None:
        self._wb.save(self.file_path)

    # ---------- BaseAdapter contract ----------

    def read_all(self, max_rows_per_sheet: int = 200) -> str:
        chunks = []
        for name in self.list_sheets():
            ws = self._sheet(name)
            chunks.append(f"--- Sheet: {name} ---")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows_per_sheet:
                    chunks.append(
                        f"... ({ws.max_row - max_rows_per_sheet} more rows truncated)"
                    )
                    break
                chunks.append(", ".join("" if v is None else str(v) for v in row))
        return "\n".join(chunks)

    def search(self, value: Any) -> list[tuple[str, str]]:
        hits: list[tuple[str, str]] = []
        for name in self.list_sheets():
            ws = self._sheet(name)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value == value:
                        hits.append((name, cell.coordinate))
        return hits

    def close(self) -> None:
        self._wb.close()

    def reload(self) -> None:
        try:
            self._wb.close()
        except Exception:
            pass
        self._wb = openpyxl.load_workbook(self.file_path, data_only=False)

    # ---------- READ ----------

    def read_cell(self, cell: str, sheet_name: Optional[str] = None) -> Any:
        return self._sheet(sheet_name)[cell].value

    def read_range(
        self, cell_range: str, sheet_name: Optional[str] = None
    ) -> list[list[Any]]:
        ws = self._sheet(sheet_name)
        return [[c.value for c in row] for row in ws[cell_range]]

    # ---------- CELL / RANGE ----------

    def write_cell(
        self, cell: str, new_value: Any, sheet_name: Optional[str] = None
    ) -> ChangePreview:
        ws = self._sheet(sheet_name)
        old_value = ws[cell].value

        def _apply():
            ws[cell] = new_value
            self._save()

        return ChangePreview(
            description="write_cell",
            source=self.file_path,
            location=f"{ws.title}!{cell}",
            old_value=old_value,
            new_value=new_value,
            _apply_fn=_apply,
        )

    def clear_range(
        self, cell_range: str, sheet_name: Optional[str] = None
    ) -> ChangePreview:
        ws = self._sheet(sheet_name)
        old = [[c.value for c in row] for row in ws[cell_range]]

        def _apply():
            for row in ws[cell_range]:
                for cell in row:
                    cell.value = None
            self._save()

        return ChangePreview(
            description="clear_range",
            source=self.file_path,
            location=f"{ws.title}!{cell_range}",
            old_value=old,
            new_value=None,
            _apply_fn=_apply,
        )

    def write_range(
        self,
        start_cell: str,
        values: list[list[Any]],
        sheet_name: Optional[str] = None,
    ) -> ChangePreview:
        if not values or not isinstance(values, list):
            raise ValueError("write_range needs a 2D values list")
        ws = self._sheet(sheet_name)
        start = ws[start_cell]
        r0, c0 = start.row, start.column

        def _apply():
            for ri, row_vals in enumerate(values):
                if not isinstance(row_vals, (list, tuple)):
                    row_vals = [row_vals]
                for ci, val in enumerate(row_vals):
                    ws.cell(row=r0 + ri, column=c0 + ci, value=val)
            self._save()

        return ChangePreview(
            description="write_range",
            source=self.file_path,
            location=f"{ws.title}!{start_cell}",
            old_value=None,
            new_value=values,
            _apply_fn=_apply,
        )

    # ---------- ROWS ----------

    def append_row(
        self, values: list[Any], sheet_name: Optional[str] = None
    ) -> ChangePreview:
        ws = self._sheet(sheet_name)
        target_row = (ws.max_row or 0) + 1

        def _apply():
            ws.append(values)
            self._save()

        return ChangePreview(
            description="append_row",
            source=self.file_path,
            location=f"{ws.title}!row {target_row}",
            old_value=None,
            new_value=values,
            _apply_fn=_apply,
        )

    def insert_row(
        self,
        row_index: int,
        values: Optional[list[Any]] = None,
        sheet_name: Optional[str] = None,
    ) -> ChangePreview:
        if row_index < 1:
            raise ValueError("row_index must be >= 1")
        ws = self._sheet(sheet_name)

        def _apply():
            ws.insert_rows(row_index)
            if values:
                for i, val in enumerate(values, start=1):
                    ws.cell(row=row_index, column=i, value=val)
            self._save()

        return ChangePreview(
            description="insert_row",
            source=self.file_path,
            location=f"{ws.title}!row {row_index}",
            old_value=None,
            new_value=values,
            _apply_fn=_apply,
        )

    def delete_row(
        self, row_index: int, sheet_name: Optional[str] = None
    ) -> ChangePreview:
        ws = self._sheet(sheet_name)
        old_values = [c.value for c in ws[row_index]]

        def _apply():
            ws.delete_rows(row_index)
            self._save()

        return ChangePreview(
            description="delete_row",
            source=self.file_path,
            location=f"{ws.title}!row {row_index}",
            old_value=old_values,
            new_value=None,
            _apply_fn=_apply,
        )

    # ---------- COLUMNS ----------

    def insert_column(
        self, column: Any, sheet_name: Optional[str] = None
    ) -> ChangePreview:
        idx = _col_to_index(column)
        ws = self._sheet(sheet_name)
        letter = get_column_letter(idx)

        def _apply():
            ws.insert_cols(idx)
            self._save()

        return ChangePreview(
            description="insert_column",
            source=self.file_path,
            location=f"{ws.title}!col {letter}",
            old_value=None,
            new_value=letter,
            _apply_fn=_apply,
        )

    def delete_column(
        self, column: Any, sheet_name: Optional[str] = None
    ) -> ChangePreview:
        idx = _col_to_index(column)
        ws = self._sheet(sheet_name)
        letter = get_column_letter(idx)
        old = [ws.cell(row=r, column=idx).value for r in range(1, (ws.max_row or 0) + 1)]

        def _apply():
            ws.delete_cols(idx)
            self._save()

        return ChangePreview(
            description="delete_column",
            source=self.file_path,
            location=f"{ws.title}!col {letter}",
            old_value=old,
            new_value=None,
            _apply_fn=_apply,
        )

    # ---------- SHEETS ----------

    def create_sheet(
        self, name: str, index: Optional[int] = None
    ) -> ChangePreview:
        name = (name or "").strip()
        if not name:
            raise ValueError("sheet name is required")
        if name in self._wb.sheetnames:
            raise ValueError(f"Sheet already exists: {name}")

        def _apply():
            if index is None:
                self._wb.create_sheet(title=name)
            else:
                self._wb.create_sheet(title=name, index=int(index))
            self._save()

        return ChangePreview(
            description="create_sheet",
            source=self.file_path,
            location=name,
            old_value=None,
            new_value=name,
            _apply_fn=_apply,
        )

    def rename_sheet(self, old_name: str, new_name: str) -> ChangePreview:
        old_name = (old_name or "").strip()
        new_name = (new_name or "").strip()
        if old_name not in self._wb.sheetnames:
            raise KeyError(f"Sheet '{old_name}' not found")
        if not new_name:
            raise ValueError("new sheet name is required")
        if new_name in self._wb.sheetnames and new_name != old_name:
            raise ValueError(f"Sheet already exists: {new_name}")

        def _apply():
            self._wb[old_name].title = new_name
            self._save()

        return ChangePreview(
            description="rename_sheet",
            source=self.file_path,
            location=old_name,
            old_value=old_name,
            new_value=new_name,
            _apply_fn=_apply,
        )

    def delete_sheet(self, name: str) -> ChangePreview:
        name = (name or "").strip()
        if name not in self._wb.sheetnames:
            raise KeyError(f"Sheet '{name}' not found")
        if len(self._wb.sheetnames) <= 1:
            raise ValueError("Cannot delete the only sheet in the workbook")

        def _apply():
            del self._wb[name]
            self._save()

        return ChangePreview(
            description="delete_sheet",
            source=self.file_path,
            location=name,
            old_value=name,
            new_value=None,
            _apply_fn=_apply,
        )

    def copy_sheet(self, name: str, new_name: str) -> ChangePreview:
        name = (name or "").strip()
        new_name = (new_name or "").strip()
        if name not in self._wb.sheetnames:
            raise KeyError(f"Sheet '{name}' not found")
        if not new_name:
            raise ValueError("new sheet name is required")
        if new_name in self._wb.sheetnames:
            raise ValueError(f"Sheet already exists: {new_name}")

        def _apply():
            src = self._wb[name]
            dst = self._wb.copy_worksheet(src)
            dst.title = new_name
            self._save()

        return ChangePreview(
            description="copy_sheet",
            source=self.file_path,
            location=f"{name} → {new_name}",
            old_value=name,
            new_value=new_name,
            _apply_fn=_apply,
        )
